# -*- coding: utf-8 -*-
"""
0608 数据集导入脚本
覆盖 t_risk_label（关键词表）和 t_case（案例表）

运行方式：
    cd C:/Users/29488/Desktop/项目
    python db/import_0608.py
"""
import os, sys
import openpyxl
import mysql.connector
from datetime import datetime, date

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
KW_FILE   = os.path.join(BASE, "0608数据集", "关键词映射关系-青少年分类0608.xlsx")
CASE_FILE = os.path.join(BASE, "0608数据集", "案例整理20260608(1).xlsx")

DB = dict(host=os.environ.get("DB_HOST", "127.0.0.1"), port=int(os.environ.get("DB_PORT", "3306")),
          user=os.environ.get("DB_USER", "root"), password=os.environ.get("DB_PASSWORD", ""),
          database=os.environ.get("DB_NAME", "youth_ideology"), charset="utf8mb4", use_pure=True)

# ── 领域名标准化（Excel 里写"XX领域"，DB 存"XX风险"）──
DOMAIN_MAP = {
    "制度认同领域": "制度认同风险", "制度认同风险": "制度认同风险",
    "历史认知领域": "历史认知风险", "历史认知风险": "历史认知风险",
    "心理韧性领域": "心理韧性风险", "心理韧性风险": "心理韧性风险",
    "认知闭合领域": "认知闭合风险", "认知闭合风险": "认知闭合风险",
    "网络素养领域": "网络素养风险", "网络素养风险": "网络素养风险",
}

# V1-V4 → 风险分值
V_SCORE = {"V1": 1.0, "V2": 2.0, "V3": 3.0, "V4": 4.0}

# V1-V4 → 对立等级（A高/B中/C低，0608删了原列，从V推导）
V_TO_OPP = {"V1": "C", "V2": "B", "V3": "A", "V4": "A"}

# 案例风险等级 V→L 映射
V_TO_L = {"V1": "L1", "V2": "L2", "V3": "L3", "V4": "L4"}


def read_sheet(path, sheet_name):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return []
    headers = [str(h).strip() if h is not None else f"col_{i}"
               for i, h in enumerate(rows[0])]
    result = []
    for row in rows[1:]:
        if all(v is None or str(v).strip() in ("", "nan") for v in row):
            continue
        result.append({headers[i]: (row[i] if i < len(row) else None)
                       for i in range(len(headers))})
    return result


def _str(val, default=""):
    return str(val).strip() if val is not None and str(val).strip() not in ("", "nan", "None") else default


def _float(val):
    try:
        return float(str(val).strip())
    except Exception:
        return None


def _date(val):
    if val is None:
        return None
    if isinstance(val, (datetime, date)):
        return val.date() if isinstance(val, datetime) else val
    s = str(val).strip()
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%Y/%m/%d", "%Y%m%d"):
        try:
            return datetime.strptime(s[:len(fmt)], fmt).date()
        except Exception:
            pass
    return None


# ══════════════════════════════════════════
#  1. 导入关键词表 t_risk_label
# ══════════════════════════════════════════
def import_keywords(cur):
    print("\n[t_risk_label] 清空旧数据...")
    cur.execute("DELETE FROM t_risk_label")

    rows = read_sheet(KW_FILE, "Sheet1")
    print(f"[t_risk_label] 读取 {len(rows)} 条关键词，开始写入...")

    # 0608 版列名（内容性质/对立等级 两列已删除，用 V 推导）
    V_COL = "群体对立风险阈值（低风险V1、中风险V2、高风险V3、极高风险V4）"

    inserted, errors = 0, []
    seen_kw = set()   # 去重：同一 Excel 里可能有重复关键词

    for i, row in enumerate(rows):
        kw = _str(row.get("关键词"))
        if not kw or kw in seen_kw:
            continue
        seen_kw.add(kw)

        tag_id = f"TAG-{i+1:05d}"

        v_raw   = _str(row.get(V_COL)).upper()
        v_score = V_SCORE.get(v_raw)          # 1.0 / 2.0 / 3.0 / 4.0 / None
        opp_lv  = V_TO_OPP.get(v_raw, "")     # A / B / C

        domain_raw = _str(row.get("关键词映射青少年意识形态领域"))
        domain     = DOMAIN_MAP.get(domain_raw, domain_raw)

        try:
            cur.execute("""
                INSERT INTO t_risk_label (
                    tag_id, keyword,
                    risk_domain_l1, risk_type_l2,
                    applicable_scenario, dimension,
                    opposition_level, content_nature, group_opposition_threshold,
                    content_risk_score, opposition_risk_score, total_risk_score,
                    l1_tag_code, l2_tag_code, l3_tag_code,
                    guidance_direction,
                    created_at, updated_at
                ) VALUES (
                    %s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,
                    NOW(), NOW()
                )
            """, (
                tag_id, kw,
                domain,
                _str(row.get("领域分类")),
                _str(row.get("对应场景")),
                _str(row.get("对应维度")),
                opp_lv or None,
                v_raw or None,          # content_nature 存原始 V 码
                v_raw or None,          # group_opposition_threshold
                v_score,                # content_risk_score（用 V 值近似）
                v_score,                # opposition_risk_score
                v_score,                # total_risk_score
                _str(row.get("L1")),
                _str(row.get("L2")),
                _str(row.get("L3")),
                _str(row.get("关键词对立概念说明")),  # 新增说明列 → guidance_direction
            ))
            inserted += 1
        except Exception as e:
            errors.append(f"[{kw}] {e}")

    print(f"[t_risk_label] 完成：写入 {inserted} 条")
    if errors:
        print(f"  错误 {len(errors)} 条：", errors[:5])
    return inserted


# ══════════════════════════════════════════
#  2. 导入案例表 t_case
# ══════════════════════════════════════════
def import_cases(cur):
    print("\n[t_case] 清空旧数据...")
    cur.execute("DELETE FROM t_case")

    rows = read_sheet(CASE_FILE, "标准案例")
    print(f"[t_case] 读取 {len(rows)} 条案例，开始写入...")

    inserted, errors = 0, []

    for row in rows:
        case_id = _str(row.get("案例 ID") or row.get("案例ID"))
        if not case_id:
            continue

        # 风险等级：V→L 映射，空值保持 NULL
        v_lv  = _str(row.get("风险等级")).upper()
        risk_level = V_TO_L.get(v_lv) or None

        try:
            cur.execute("""
                INSERT INTO t_case (
                    case_id, case_name,
                    risk_domain_l1, risk_type_l2,
                    platform, risk_level,
                    incident_time, source_url,
                    spread_path, spread_mechanism, interaction_feature,
                    representative_text,
                    disposal_strategy_platform,
                    disposal_strategy_6_12,
                    disposal_strategy_13_15,
                    disposal_strategy_16_18,
                    review_conclusion,
                    created_at, updated_at
                ) VALUES (
                    %s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,
                    NOW(), NOW()
                )
            """, (
                case_id,
                _str(row.get("案例标题"))[:290],   # VARCHAR(300) 限制
                _str(row.get("一级风险域")),
                _str(row.get("二级标签")),
                _str(row.get("传播平台")),
                risk_level,
                _date(row.get("发生时间")),
                _str(row.get("数据来源（网页或图片链接）")),
                _str(row.get("传播路径")),
                # spread_mechanism 用"传播效果"+"互动特征"拼一下（案例整理没有专门的传播机制列）
                " | ".join(filter(None, [
                    _str(row.get("传播效果")),
                    _str(row.get("互动特征")),
                ])) or None,
                _str(row.get("互动特征")),
                _str(row.get("内容摘要")),
                _str(row.get("处置方式（平台）")),
                _str(row.get("处置方式（6-12岁认知启蒙期青少年引导）")),
                _str(row.get("处置方式（13-15岁价值观形成期青少年引导）")),
                _str(row.get("处置方式（16-18岁理性思辨期青少年引导）")),
                _str(row.get("复盘结论")),
            ))
            inserted += 1
        except Exception as e:
            errors.append(f"[{case_id}] {e}")

    print(f"[t_case] 完成：写入 {inserted} 条")
    if errors:
        print(f"  错误 {len(errors)} 条：", errors[:5])
    return inserted


# ══════════════════════════════════════════
#  主流程
# ══════════════════════════════════════════
def main():
    print("=" * 50)
    print("0608 数据集导入")
    print("=" * 50)

    conn = mysql.connector.connect(**DB)
    cur  = conn.cursor()
    cur.execute("SET NAMES utf8mb4")
    cur.execute("SET foreign_key_checks = 0")

    try:
        kw_cnt   = import_keywords(cur)
        case_cnt = import_cases(cur)
        conn.commit()

        # 验证
        print("\n── 导入结果验证 ──")
        for tbl in ("t_risk_label", "t_case"):
            cur.execute(f"SELECT COUNT(*) FROM `{tbl}`")
            print(f"  {tbl}: {cur.fetchone()[0]} 行")

        # 案例风险等级分布
        cur.execute("SELECT risk_level, COUNT(*) FROM t_case GROUP BY risk_level")
        print("  t_case 风险等级分布:", dict(cur.fetchall()))

        # 关键词风险域分布
        cur.execute("SELECT risk_domain_l1, COUNT(*) FROM t_risk_label GROUP BY risk_domain_l1 ORDER BY COUNT(*) DESC")
        print("  t_risk_label 风险域分布:", dict(cur.fetchall()))

    except Exception as e:
        conn.rollback()
        print(f"\n[严重错误] {e}，已回滚")
        raise
    finally:
        cur.execute("SET foreign_key_checks = 1")
        cur.close()
        conn.close()

    print("\n完成！重启后端服务让新数据生效。")


if __name__ == "__main__":
    main()
