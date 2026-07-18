# -*- coding: utf-8 -*-
"""
补充缺失关键词到 t_risk_label
来源：0529数据集/关键词映射关系-青少年分类0529.xlsx  Sheet1（214 条）
已存在的关键词跳过，只插入缺失的。
"""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl
import mysql.connector


def read_sheet(file_path, sheet_name=None):
    """读取 Excel sheet，返回 list[dict]，key 为列头。"""
    wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
    ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return []
    headers = [str(h).strip() if h is not None else f"col_{i}"
               for i, h in enumerate(rows[0])]
    result = []
    for row in rows[1:]:
        if all(v is None or str(v).strip() in ("", "nan") for v in row):
            continue
        result.append({headers[i]: (row[i] if i < len(row) else None)
                       for i in range(len(headers))})
    return result

# ── 数据库连接 ──
DB = dict(host=os.environ.get("DB_HOST", "127.0.0.1"), port=int(os.environ.get("DB_PORT", "3306")),
          user=os.environ.get("DB_USER", "root"), password=os.environ.get("DB_PASSWORD", ""),
          database=os.environ.get("DB_NAME", "youth_ideology"), charset="utf8mb4", use_pure=True)

# ── 领域名称标准化（Excel 里写的是"XX领域"，DB 里存的是"XX风险"）──
DOMAIN_MAP = {
    "制度认同领域": "制度认同风险",
    "历史认知领域": "历史认知风险",
    "心理韧性领域": "心理韧性风险",
    "认知闭合领域": "认知闭合风险",
    "网络素养领域": "网络素养风险",
    # 部分条目直接就写了"XX风险"，保持原样
    "制度认同风险": "制度认同风险",
    "历史认知风险": "历史认知风险",
    "心理韧性风险": "心理韧性风险",
    "认知闭合风险": "认知闭合风险",
    "网络素养风险": "网络素养风险",
}

# C1~C4 / V1~V4 → 数值分
LEVEL_SCORE = {"C1": 1.0, "C2": 2.0, "C3": 3.0, "C4": 4.0,
               "V1": 1.0, "V2": 2.0, "V3": 3.0, "V4": 4.0}


def _score(val):
    """把 C1/V1 等字符串转成 float，解析失败返回 None。"""
    v = str(val or "").strip().upper()
    return LEVEL_SCORE.get(v)


def run():
    conn = mysql.connector.connect(**DB)
    cur = conn.cursor(dictionary=True)

    # ── 1. 读取现有关键词，建去重集合 ──
    cur.execute("SELECT keyword FROM t_risk_label")
    existing_kw = {r["keyword"] for r in cur.fetchall()}
    print(f"现有关键词：{len(existing_kw)} 条")

    # ── 2. 读取现有最大 tag_id 编号 ──
    cur.execute("SELECT MAX(tag_id) AS mx FROM t_risk_label")
    mx = cur.fetchone()["mx"] or "TAG-00000"
    next_no = int(mx.replace("TAG-", "")) + 1

    # ── 3. 读取 Excel ──
    xlsx = os.path.join(
        os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
        "0529数据集", "关键词映射关系-青少年分类0529.xlsx"
    )
    rows = read_sheet(xlsx, "Sheet1")
    print(f"Excel 共 {len(rows)} 条关键词")

    # ── 4. 逐行处理 ──
    inserted, skipped, errors = 0, 0, []

    for row in rows:
        kw = str(row.get("关键词") or "").strip()
        if not kw or kw == "关键词":
            continue

        if kw in existing_kw:
            skipped += 1
            continue

        tag_id = f"TAG-{next_no:05d}"
        next_no += 1

        # 分值
        c_val = _score(row.get("内容性质（低风险C1、中风险C2、高风险C3、极高风险C4）"))
        v_val = _score(row.get("群体对立风险阈值（低风险V1、中风险V2、高风险V3、极高风险V4）"))
        # total = 两项均有时取均值，否则取有值的那个
        if c_val is not None and v_val is not None:
            total = round((c_val + v_val) / 2, 2)
        else:
            total = c_val or v_val

        # 风险域
        raw_domain = str(row.get("关键词映射青少年意识形态领域") or "").strip()
        domain = DOMAIN_MAP.get(raw_domain, raw_domain) or ""

        try:
            cur.execute("""
                INSERT INTO t_risk_label (
                    tag_id, keyword,
                    risk_domain_l1, risk_type_l2,
                    applicable_scenario, dimension,
                    opposition_level, content_nature, group_opposition_threshold,
                    content_risk_score, opposition_risk_score, total_risk_score,
                    l1_tag_code, l2_tag_code, l3_tag_code,
                    created_at, updated_at
                ) VALUES (
                    %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s,
                    NOW(), NOW()
                )
            """, (
                tag_id,
                kw,
                domain,
                str(row.get("领域分类") or "").strip(),
                str(row.get("对应场景") or "").strip(),
                str(row.get("对应维度") or "").strip(),
                str(row.get("对立等级（A高对立、B中对立、C低对立）") or "").strip(),
                str(row.get("内容性质（低风险C1、中风险C2、高风险C3、极高风险C4）") or "").strip(),
                str(row.get("群体对立风险阈值（低风险V1、中风险V2、高风险V3、极高风险V4）") or "").strip(),
                c_val,
                v_val,
                total,
                str(row.get("L1") or "").strip(),
                str(row.get("L2") or "").strip(),
                str(row.get("L3") or "").strip(),
            ))
            existing_kw.add(kw)   # 防止同一 Excel 里有重复行
            inserted += 1
        except Exception as e:
            errors.append(f"[{kw}] {e}")

    conn.commit()
    cur.close()
    conn.close()

    print(f"\n完成：插入 {inserted} 条，跳过（已存在）{skipped} 条")
    if errors:
        print(f"错误 {len(errors)} 条：")
        for e in errors[:10]:
            print(" ", e)


if __name__ == "__main__":
    run()
